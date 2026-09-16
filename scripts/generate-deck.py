"""
Builds the MediBloom hackathon pitch deck.

    python scripts/generate-deck.py     ->  MediBloom-Pitch.pptx

Each slide is authored as HTML, rendered by headless Chrome at 2560x1440 and
dropped full-bleed into a 16:9 PPTX. That is the same production route as the
original deck, which is why the two look like the same product: DM Serif
Display for headings, Nunito for everything else, and the app's own palette.

Phone screenshots come from the signed release APK. If they are missing the
deck still builds, with those frames left empty.
"""

import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.environ.get(
    "DECK_SHOTS",
    r"C:\Users\saish\AppData\Local\Temp\claude\C--Users-saish"
    r"\d1f4b857-5311-4f0a-b7a2-07f6fc03db46\scratchpad\shots\deck",
)
OUT_PPTX = os.path.join(ROOT, "MediBloom-Pitch.pptx")
RENDER_DIR = os.path.join(tempfile.gettempdir(), "medibloom-deck-render")

TOTAL = 8  # numbered slides; anything after is appendix

CHROME_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]

# --------------------------------------------------------------------------- #
#  The flower mark, same geometry as the Logo component in the app.
# --------------------------------------------------------------------------- #
PETALS = [(0, "#E85D8A", .9), (60, "#F4A6C1", .85), (120, "#8B5FBF", .9),
          (180, "#C9B6E4", .85), (240, "#D4A574", .85), (300, "#E85D8A", .8)]


def flower(px):
    petals = "".join(
        f'<ellipse cx="0" cy="-10" rx="5.5" ry="10" fill="{c}" opacity="{o}" '
        f'transform="rotate({r})"/>' for r, c, o in PETALS)
    return (
        f'<svg width="{px}" height="{px}" viewBox="0 0 40 40" '
        f'style="display:block;flex:none">'
        f'<g transform="translate(20,20)">{petals}'
        f'<circle cx="0" cy="0" r="6.5" fill="#fff"/>'
        f'<rect x="-4" y="-1.4" width="8" height="2.8" rx="1.4" fill="#8B5FBF" '
        f'transform="rotate(45)"/></g></svg>')


CSS = """
  @import url('https://fonts.googleapis.com/css2?family=DM+Serif+Display:ital@0;1&family=Nunito:wght@400;600;700;800&display=swap');

  *, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }

  :root {
    --ink:#2D2438; --ink-soft:#6B6178; --ink-faint:#9A90A6;
    --rose:#E85D8A; --rose-deep:#C63C68; --rose-soft:#FDEBF1;
    --violet:#8B5FBF; --violet-deep:#6E45A3; --violet-soft:#F1EBF8;
    --sage:#3F8F6B; --sage-soft:#E7F3ED;
    --amber:#9A6F2E; --amber-soft:#FAF1E2;
    --coral:#D94F4F; --coral-soft:#FDECEC;
    --card:#FFFFFF; --line:rgba(45,36,56,.07);
    --serif:'DM Serif Display', Georgia, serif;
    --sans:'Nunito', 'Segoe UI', system-ui, sans-serif;
  }

  body {
    width:1280px; height:720px; overflow:hidden;
    font-family:var(--sans); color:var(--ink);
    background:#FDF8F6;
    -webkit-font-smoothing:antialiased;
  }

  /* The soft blooms in the corners, same as the original deck. */
  .slide {
    position:relative; width:1280px; height:720px; padding:46px 64px;
    display:flex; flex-direction:column;
    background:
      radial-gradient(760px 520px at 100% 0%,   rgba(232,93,138,.10), transparent 62%),
      radial-gradient(620px 460px at 0% 100%,   rgba(139,95,191,.09), transparent 60%),
      radial-gradient(520px 380px at 12% 6%,    rgba(212,165,116,.06), transparent 60%),
      #FDF8F6;
  }

  .brandbar { display:flex; align-items:center; justify-content:space-between; margin-bottom:26px; }
  .brand { display:flex; align-items:center; gap:9px; }
  .brand span { font-weight:800; font-size:16px; color:var(--violet); letter-spacing:.01em; }
  .counter { font-weight:800; font-size:14px; color:#C9BFD4; letter-spacing:.16em; }
  .badge {
    font-weight:800; font-size:12px; letter-spacing:.14em; color:var(--violet-deep);
    background:var(--violet-soft); padding:7px 16px; border-radius:999px;
  }

  h1 { font-family:var(--serif); font-size:56px; line-height:1.04; letter-spacing:-.01em; font-weight:400; }
  .sub { font-size:18px; color:var(--ink-soft); margin-top:9px; max-width:960px; line-height:1.45; }

  .card {
    background:var(--card); border:1px solid var(--line); border-radius:18px;
    box-shadow:0 1px 3px rgba(45,36,56,.04), 0 8px 24px rgba(45,36,56,.035);
  }

  .icon {
    width:44px; height:44px; border-radius:13px;
    display:grid; place-items:center; font-size:20px; flex:none;
  }

  .eyebrow { font-weight:800; font-size:12px; letter-spacing:.15em; text-transform:uppercase; }
  .quote { font-family:var(--serif); font-style:italic; font-size:29px; color:var(--violet-deep); text-align:center; }

  .flow { display:flex; align-items:center; gap:10px; }
  .step {
    flex:1; text-align:center; font-weight:700; font-size:15px;
    padding:15px 8px; border-radius:14px; background:#fff;
    border:1px solid var(--line); box-shadow:0 1px 3px rgba(45,36,56,.04);
  }
  .step.end {
    background:linear-gradient(100deg, var(--rose), var(--violet));
    color:#fff; border:none;
  }
  .arrow { color:#CDBFDA; font-size:19px; font-weight:800; flex:none; }

  .grid2 { display:grid; grid-template-columns:1fr 1fr; gap:14px; }
  .grid3 { display:grid; grid-template-columns:repeat(3,1fr); gap:14px; }
  .grid4 { display:grid; grid-template-columns:repeat(4,1fr); gap:12px; }

  .row {
    display:flex; align-items:center; gap:16px;
    background:#fff; border:1px solid var(--line); border-radius:13px;
    padding:13px 18px; box-shadow:0 1px 2px rgba(45,36,56,.03);
  }

  .stat { text-align:left; }
  .stat .n { font-family:var(--serif); font-size:40px; line-height:1; color:var(--rose-deep); }
  .stat .l { font-size:13.5px; color:var(--ink-soft); margin-top:7px; font-weight:600; }

  .phone { border-radius:20px; background:var(--ink); padding:5px; flex:none; box-shadow:0 10px 30px rgba(45,36,56,.16); }
  .phone img { display:block; border-radius:16px; }

  .note { font-size:13.5px; color:var(--ink-faint); line-height:1.5; }
  .fill { flex:1; }
  b, strong { font-weight:800; }
"""


def head(counter=None, badge=None):
    right = ""
    if badge:
        right = f'<div class="badge">{badge}</div>'
    elif counter:
        right = f'<div class="counter">{counter:02d} / {TOTAL:02d}</div>'
    return (f'<div class="brandbar"><div class="brand">{flower(26)}'
            f'<span>MediBloom</span></div>{right}</div>')


def title_block(h, sub=None):
    s = f'<p class="sub">{sub}</p>' if sub else ""
    return f"<h1>{h}</h1>{s}"


def img_tag(name, width):
    """A phone screenshot at a given CSS width, or an empty frame if absent."""
    path = os.path.join(SHOTS, name)
    h = round(width * 2400 / 1080)
    if not os.path.exists(path):
        return (f'<div class="phone"><div style="width:{width}px;height:{h}px;'
                f'border-radius:16px;background:#3A3145"></div></div>')
    url = "file:///" + path.replace("\\", "/")
    return f'<div class="phone"><img src="{url}" width="{width}" height="{h}"></div>'


# --------------------------------------------------------------------------- #
#  Slides
# --------------------------------------------------------------------------- #

def s01_title():
    return f"""
<div class="slide" style="align-items:center;justify-content:center;text-align:center;padding-top:0">
  <div style="margin-bottom:22px">{flower(104)}</div>
  <div style="font-family:var(--serif);font-size:88px;line-height:1">MediBloom</div>
  <div style="font-weight:800;font-size:25px;color:var(--violet-deep);margin-top:16px">
    Your medicines, remembered. Your safety, checked.
  </div>
  <p style="font-size:18px;color:var(--ink-soft);margin-top:14px;max-width:760px;line-height:1.5">
    An offline-first Android app that checks every pair of your medicines for dangerous
    interactions, reminds you when doses are due, and explains the risk in plain language.
  </p>
  <div style="display:flex;gap:12px;margin-top:30px">
    <div style="background:var(--violet-soft);color:var(--violet-deep);font-weight:800;font-size:15px;padding:11px 24px;border-radius:999px">100% offline-first</div>
    <div style="background:var(--sage-soft);color:var(--sage);font-weight:800;font-size:15px;padding:11px 24px;border-radius:999px">&#8377;0 to run, forever</div>
    <div style="background:var(--rose-soft);color:var(--rose-deep);font-weight:800;font-size:15px;padding:11px 24px;border-radius:999px">Private by default</div>
  </div>
  <div style="position:absolute;bottom:44px;font-size:14px;color:var(--ink-faint);font-weight:600">
    [Your Team Name] &nbsp;&middot;&nbsp; Healthtech &amp; Wellness &nbsp;&middot;&nbsp; [Hackathon Name, 2026]
  </div>
</div>"""


def s02_problem():
    cards = [
        ("&#9888;", "var(--coral-soft)", "Risky combinations hide in plain sight",
         "Nobody holds the full list. A painkiller bought at the medical shop never "
         "reaches the doctor who prescribed the rest."),
        ("&#128276;", "var(--violet-soft)", "Reminders that don't stick",
         "One phone alarm gets snoozed, silenced and forgotten. Evening doses slip the most."),
        ("&#128274;", "var(--amber-soft)", "Help lives in the wrong place",
         "Interaction checkers exist &mdash; online, in English, built for pharmacists. "
         "Not woven into anyone's daily routine."),
    ]
    html = "".join(
        f'<div class="card" style="padding:24px"><div class="icon" style="background:{bg}">{ic}</div>'
        f'<div style="font-weight:800;font-size:18px;margin-top:16px">{t}</div>'
        f'<div style="font-size:14.5px;color:var(--ink-soft);margin-top:9px;line-height:1.5">{b}</div></div>'
        for ic, bg, t, b in cards)
    return f"""
<div class="slide">
  {head(2)}
  {title_block("The problem", "Taking several medicines safely is harder than it should be.")}
  <div class="grid3" style="margin-top:26px">{html}</div>
  <div class="card" style="margin-top:16px;padding:17px 24px;display:flex;align-items:center;gap:14px">
    <span style="color:var(--amber);font-size:19px">&#9888;</span>
    <div style="font-weight:700;font-size:15.5px">
      The WHO puts adherence to long-term therapy at roughly <b>50%</b>. Taking five or more
      medicines &mdash; normal for an older adult &mdash; roughly <b>doubles</b> the risk of an adverse event.
    </div>
  </div>
  <div class="fill"></div>
  <div class="quote">&ldquo;What if checking your medicines was as easy as taking them?&rdquo;</div>
  <div class="fill"></div>
</div>"""


def s03_solution():
    feats = [
        ("&#128138;", "var(--rose-soft)", "var(--rose-deep)", "Medication tracking",
         "Every dose, every time, in one place"),
        ("&#9888;", "var(--coral-soft)", "var(--coral)", "Interaction safety",
         "Checked the moment you add a medicine"),
        ("&#128276;", "var(--violet-soft)", "var(--violet-deep)", "Reminders that act",
         "Mark taken, snooze or skip from the notification"),
        ("&#127807;", "var(--sage-soft)", "var(--sage)", "Wellness insights",
         "Plain-language patterns from your own history"),
    ]
    cards = "".join(
        f'<div style="background:{bg};border-radius:15px;padding:18px 20px;display:flex;align-items:center;gap:15px">'
        f'<div class="icon" style="background:#fff;color:{fg}">{ic}</div><div>'
        f'<div style="font-weight:800;font-size:17px">{t}</div>'
        f'<div style="font-size:14px;color:{fg};margin-top:3px">{b}</div></div></div>'
        for ic, bg, fg, t, b in feats)
    return f"""
<div class="slide">
  {head(3)}
  {title_block("Our solution", "One app: log your medicines, and MediBloom takes care of the rest.")}
  <div class="flow" style="margin-top:26px">
    <div class="step">You log your medicines</div><div class="arrow">&rsaquo;</div>
    <div class="step">MediBloom checks every pair</div><div class="arrow">&rsaquo;</div>
    <div class="step">It reminds &amp; explains</div><div class="arrow">&rsaquo;</div>
    <div class="step end">Safer, steadier days</div>
  </div>
  <div class="eyebrow" style="color:var(--ink);margin:28px 0 12px">MediBloom brings together</div>
  <div class="grid2">{cards}</div>
  <div class="fill"></div>
  <div class="card" style="padding:18px 24px;display:flex;align-items:center;gap:16px">
    <div style="font-weight:800;font-size:16px;color:var(--rose-deep);white-space:nowrap">Smart Swap</div>
    <div style="font-size:15px;color:var(--ink-soft)">
      We don't just warn. Where a safer alternative exists we name it &mdash; and we swap the
      painkiller, never the medicine keeping you alive.
    </div>
  </div>
</div>"""


def s04_features():
    rows = [
        ("Interaction checking", "134 rules &middot; 130 generics &middot; 398 brand names",
         "Every pair re-checked instantly, offline. Indian brands resolve: Meftal, Shelcal, Thyronorm."),
        ("Prescription scanning", "Photo or PDF &middot; on-device OCR",
         "Google ML Kit reads it on the phone. It shows what it saw and asks before saving anything."),
        ("Reminders", "Exact alarms &middot; action buttons",
         "Mark taken, snooze or skip from the lock screen. Survives reboot."),
        ("Caregiver alerts", "Automatic email",
         "A dose left unmarked past your limit emails a family member, with no server involved."),
        ("Ask MediBloom", "Offline answers + optional AI",
         "&ldquo;What did I have today?&rdquo; answers instantly from the phone. Open-ended questions go to a model."),
        ("Wellness dashboard", "Streaks &middot; 4-week calendar &middot; trends",
         "Tap any day to see that day's doses. Week-on-week comparison, per-medicine breakdown."),
    ]
    html = "".join(
        f'<div class="row"><div style="width:218px;flex:none">'
        f'<div style="font-weight:800;font-size:16px">{t}</div>'
        f'<div style="font-size:12.5px;color:var(--violet-deep);font-weight:700;margin-top:2px">{s}</div></div>'
        f'<div style="font-size:14.5px;color:var(--ink-soft);line-height:1.45">{b}</div></div>'
        for t, s, b in rows)
    return f"""
<div class="slide">
  {head(4)}
  {title_block("Key features", "Everything below is built, signed and running on a real phone today.")}
  <div style="display:flex;flex-direction:column;gap:9px;margin-top:24px">{html}</div>
  <div class="fill"></div>
  <div class="note">Plus larger-text mode, voice read-aloud, &ge;44px touch targets and plain
  language throughout &mdash; accessibility is the brief, not a checkbox.</div>
</div>"""


def s05_demo():
    return f"""
<div class="slide">
  {head(5)}
  {title_block("See it working", "Lakshmi, 68, Chennai &mdash; a medicine list from a real Indian kitchen shelf.")}
  <div style="display:flex;gap:26px;margin-top:22px;align-items:flex-start">
    <div style="flex:1;padding-top:4px">
      <div style="font-size:15.5px;color:var(--ink-soft);line-height:1.55">
        Warfarin since a heart valve replacement. Thyronorm and Shelcal, like half the older
        women in the country. Glycomet for her diabetes. Nurokind for the tingling in her feet.
      </div>
      <div style="background:var(--coral-soft);border-radius:15px;padding:18px 20px;margin-top:16px">
        <div style="font-weight:800;font-size:16.5px;color:var(--coral)">
          And Meftal, bought over the counter for knee pain.
        </div>
        <div style="font-size:14.5px;color:var(--ink-soft);margin-top:7px;line-height:1.5">
          Severe bleeding risk with warfarin. No pharmacist ever saw both boxes.
          MediBloom caught it instantly, offline.
        </div>
      </div>
      <div class="card" style="padding:16px 20px;margin-top:13px">
        <div class="eyebrow" style="color:var(--violet-deep);margin-bottom:9px">The other two it finds</div>
        <div style="font-size:14px;line-height:1.7">
          <b>Thyronorm + Shelcal</b> <span style="color:var(--ink-soft)">&mdash; calcium blocks the thyroid tablet. Real, common, rarely explained.</span><br>
          <b>Nurokind + Glycomet</b> <span style="color:var(--ink-soft)">&mdash; deliberate and fine. The app has to be able to say that too.</span>
        </div>
      </div>
    </div>
    {img_tag("02-interactions.png", 224)}
    {img_tag("06-review.png", 224)}
  </div>
</div>"""


def s06_how():
    return f"""
<div class="slide">
  {head(6)}
  {title_block("How it works", "Rules answer what they can. A model answers the rest.")}
  <div style="display:flex;gap:24px;margin-top:22px;align-items:flex-start">
    <div style="flex:1">
      <div style="display:flex;gap:13px">
        <div style="flex:1;background:var(--rose-soft);border-radius:15px;padding:18px">
          <div style="font-family:var(--serif);font-size:22px;color:var(--rose-deep)">On the phone</div>
          <div style="font-weight:700;font-size:12.5px;color:var(--ink-faint);margin-top:3px">instant &middot; works in airplane mode</div>
          <div style="font-size:14px;color:var(--ink-soft);margin-top:12px;line-height:1.75">
            Interaction checking<br>What you took, what's left<br>Adherence &amp; streaks<br>Prescription parsing
          </div>
        </div>
        <div style="flex:1;background:var(--violet-soft);border-radius:15px;padding:18px">
          <div style="font-family:var(--serif);font-size:22px;color:var(--violet-deep)">Cloud assist</div>
          <div style="font-weight:700;font-size:12.5px;color:var(--ink-faint);margin-top:3px">optional &middot; falls back on failure</div>
          <div style="font-size:14px;color:var(--ink-soft);margin-top:12px;line-height:1.75">
            Open-ended questions<br>&ldquo;My knee hurts, what can I take&rdquo;<br>Plain-language explanations<br>OCR second opinion
          </div>
        </div>
      </div>
      <div class="card" style="padding:16px 20px;margin-top:13px">
        <div style="font-size:14.5px;color:var(--ink-soft);line-height:1.55">
          If the model is slow, wrong, rate-limited or unreachable,
          <b style="color:var(--ink)">the rule engine answers instead and the app says so.</b>
        </div>
        <div style="font-weight:800;font-size:14.5px;color:var(--sage);margin-top:7px">
          The demo cannot be broken by the venue wifi.
        </div>
      </div>
      <div class="note" style="margin-top:13px">
        We split it this way after measuring: the model added ~10 seconds and once rendered a
        <b>skipped</b> dose as &ldquo;you took it&rdquo;. Factual questions never leave the phone.
      </div>
    </div>
    {img_tag("04-chat.png", 232)}
  </div>
</div>"""


def s07_stack():
    groups = [
        ("App", "var(--rose-deep)",
         "React Native 0.86 &middot; Expo SDK 57<br>TypeScript, strict mode<br>"
         "React Navigation 7 &middot; react-native-svg"),
        ("On device", "var(--violet-deep)",
         "expo-sqlite &mdash; the only place data lives<br>expo-notifications &mdash; exact alarms<br>"
         "Google ML Kit &mdash; offline OCR<br>Custom Kotlin module for PDF rendering"),
        ("Optional network", "var(--sage)",
         "Groq &mdash; openai/gpt-oss-120b, free tier<br>Email&#8201;JS &mdash; 200 emails a month, free<br>"
         "or the user's own Google Apps Script"),
        ("Build &amp; test", "var(--amber)",
         "Signed Gradle release APK<br>Jest &mdash; 221 logic tests<br>adb-driven UI testing on device"),
    ]
    cells = "".join(
        f'<div><div style="font-family:var(--serif);font-size:21px;color:{fg}">{t}</div>'
        f'<div style="font-size:13.5px;color:var(--ink-soft);margin-top:9px;line-height:1.75">{b}</div></div>'
        for t, fg, b in groups)
    costs = [("Servers", "none &mdash; no backend"), ("Database", "SQLite on the phone"),
             ("Interaction data", "bundled in the APK"), ("OCR", "on-device, free"),
             ("AI", "free tier, optional"), ("Email", "free tier, optional")]
    rows = "".join(
        f'<div style="display:flex;justify-content:space-between;gap:12px;font-size:13px;padding:5px 0">'
        f'<span style="font-weight:800">{k}</span>'
        f'<span style="color:var(--ink-soft);text-align:right">{v}</span></div>'
        for k, v in costs)
    stats = "".join(
        f'<div class="card" style="padding:14px 16px"><div class="stat">'
        f'<div class="n">{n}</div><div class="l">{l}</div></div></div>'
        for n, l in [("221", "tests passing"), ("134", "interaction rules"),
                     ("12", "screens"), ("0", "crashes in logcat")])
    return f"""
<div class="slide">
  {head(7)}
  {title_block("Built &amp; validated", "Every piece chosen so it costs nothing to run &mdash; and proven on a real device.")}
  <div style="display:flex;gap:24px;margin-top:22px">
    <div style="flex:1.55">
      <div class="grid2" style="gap:20px 26px">{cells}</div>
      <div class="grid4" style="margin-top:24px">{stats}</div>
    </div>
    <div style="width:290px;flex:none;background:var(--sage-soft);border-radius:17px;padding:22px">
      <div style="font-family:var(--serif);font-size:46px;color:var(--sage);line-height:1">&#8377;0</div>
      <div style="font-weight:800;font-size:13.5px;color:var(--ink-soft);margin-top:4px">to run, per user, forever</div>
      <div style="margin-top:14px;border-top:1px solid rgba(63,143,107,.18);padding-top:8px">{rows}</div>
    </div>
  </div>
  <div class="fill"></div>
  <div class="note">No login, no account, no analytics, no tracking. One install, one person.</div>
</div>"""


def s08_future():
    near = [
        ("Pharmacist review of all 134 rules",
         "The single thing standing between this and a real patient. Already scoped."),
        ("Tamil, Hindi and Telugu",
         "The UI is already written in plain language and the layout tolerates longer strings."),
        ("Shared caregiver dashboard",
         "A read-only weekly view a daughter abroad can open &mdash; still no account for the patient."),
    ]
    later = [
        ("On-device small language model",
         "Replaces cloud assist entirely with a quantised model, so open-ended answers work offline too."),
        ("Pharmacy &amp; ABDM integration",
         "Pull the dispensed list straight from the chemist or the national health record, so nothing is missed."),
        ("Refill prediction &amp; smart scheduling",
         "Warn before the strip runs out, and suggest dose timings that avoid the clashes we detect."),
    ]

    def block(items, colour):
        return "".join(
            f'<div style="display:flex;gap:13px;align-items:flex-start">'
            f'<div style="width:9px;height:9px;border-radius:50%;background:{colour};margin-top:7px;flex:none"></div>'
            f'<div><div style="font-weight:800;font-size:15.5px">{t}</div>'
            f'<div style="font-size:13.5px;color:var(--ink-soft);margin-top:3px;line-height:1.45">{b}</div></div></div>'
            for t, b in items)

    return f"""
<div class="slide">
  {head(8)}
  {title_block("Future enhancements", "What we build next, and what it would take to put this in a real clinic.")}
  <div style="display:flex;gap:16px;margin-top:24px">
    <div class="card" style="flex:1;padding:22px">
      <div class="eyebrow" style="color:var(--rose-deep)">Next 3 months</div>
      <div style="display:flex;flex-direction:column;gap:15px;margin-top:15px">{block(near, "var(--rose)")}</div>
    </div>
    <div class="card" style="flex:1;padding:22px">
      <div class="eyebrow" style="color:var(--violet-deep)">Beyond</div>
      <div style="display:flex;flex-direction:column;gap:15px;margin-top:15px">{block(later, "var(--violet)")}</div>
    </div>
  </div>
  <div style="display:flex;gap:14px;margin-top:16px">
    <div style="flex:1;background:var(--amber-soft);border-radius:15px;padding:17px 20px">
      <div class="eyebrow" style="color:var(--amber);margin-bottom:8px">What we'd say before you ask</div>
      <div style="font-size:13.5px;color:var(--ink-soft);line-height:1.6">
        The dataset is hand-built, not certified &mdash; it needs pharmacist review.
        Handwriting OCR is unreliable, so we ask rather than guess.
        Some Android makers throttle background alarms.
      </div>
    </div>
    <div style="flex:1;background:var(--sage-soft);border-radius:15px;padding:17px 20px">
      <div class="eyebrow" style="color:var(--sage);margin-bottom:8px">Why it can scale</div>
      <div style="font-size:13.5px;color:var(--ink-soft);line-height:1.6">
        There is nothing to scale. A million users is a million phones doing their own work.
        Our cost stays at zero, and health data that never leaves the device cannot be breached.
      </div>
    </div>
  </div>
</div>"""


def s09_close():
    return f"""
<div class="slide" style="align-items:center;justify-content:center;text-align:center">
  <div style="margin-bottom:20px">{flower(74)}</div>
  <div style="font-family:var(--serif);font-size:48px;line-height:1.15;max-width:900px">
    Turn off the wifi and mobile data,<br>and do the whole demo again.
  </div>
  <div style="font-family:var(--serif);font-size:48px;color:var(--rose-deep);margin-top:6px">
    Nothing changes.
  </div>
  <p style="font-size:17px;color:var(--ink-soft);margin-top:26px;max-width:720px;line-height:1.5">
    MediBloom is built, signed and installed. 221 tests passing, 134 interaction rules,
    zero rupees to run &mdash; and it works in the places that need it most.
  </p>
  <div style="display:flex;gap:12px;margin-top:28px">
    <div style="background:#fff;border:1px solid var(--line);font-weight:800;font-size:15px;padding:11px 26px;border-radius:999px">Thank you</div>
    <div style="background:var(--violet-soft);color:var(--violet-deep);font-weight:800;font-size:15px;padding:11px 26px;border-radius:999px">Questions?</div>
  </div>
</div>"""


def a10_elderly():
    pairs = [
        ("Small text is hard to read", "Larger-text mode scales the whole type system"),
        ("Tiny buttons are hard to tap", "Every target &ge;44px; Mark taken / Snooze / Skip in one tap"),
        ("A spinner is impossible with shaky hands", "Time picker built from big + and &minus; buttons"),
        ("Not sure what a notification means", "Optional voice read-aloud says the medicine and dose"),
        ("Typing medicine names is tedious", "Brand names resolve &mdash; type &ldquo;Shelcal&rdquo;, not &ldquo;calcium carbonate&rdquo;"),
        ("Living alone, a missed dose goes unnoticed", "Caregiver alert emails a family member automatically"),
        ("Unfamiliar with apps and medical jargon", "Plain language, single-screen flows, no account or password"),
    ]
    rows = "".join(
        f'<div class="row" style="padding:11px 18px">'
        f'<div style="width:400px;flex:none;font-weight:700;font-size:14.5px;color:var(--coral)">{a}</div>'
        f'<div style="color:#CDBFDA;font-size:16px;flex:none">&rarr;</div>'
        f'<div style="font-weight:700;font-size:14.5px">{b}</div></div>'
        for a, b in pairs)
    return f"""
<div class="slide">
  {head(badge="APPENDIX &middot; ACCESSIBILITY")}
  {title_block("Designed with elderly users in mind", "Every challenge below maps to a specific feature that already exists.")}
  <div style="display:flex;flex-direction:column;gap:8px;margin-top:24px">{rows}</div>
</div>"""


def a11_testing():
    checks = [
        "Reminder fires on time; Mark taken writes to the database and Home updates",
        "Warfarin + Meftal flagged severe before the medicine is even saved",
        "Sample prescription parsed; all six Indian brands resolved to generics",
        "A deliberate misread is flagged &ldquo;Please check this&rdquo;, never guessed",
        "Cloud assist answered a free-form question in 13 seconds, safely",
        "Invalid API key &rarr; chat still answered on-device and explained why",
        "Clean logcat across every screen: no crashes, no JS errors",
        "Typing verified character-exact after fixing a dropped-keystroke bug",
    ]
    items = "".join(
        f'<div style="display:flex;gap:11px;align-items:flex-start">'
        f'<div style="color:var(--sage);font-weight:800;font-size:15px;flex:none">&check;</div>'
        f'<div style="font-size:14px;color:var(--ink-soft);line-height:1.45">{c}</div></div>'
        for c in checks)
    return f"""
<div class="slide">
  {head(badge="APPENDIX &middot; VALIDATION")}
  {title_block("How we know it works", "221 automated tests, plus the signed APK driven on a device over adb.")}
  <div class="grid2" style="margin-top:24px;gap:13px 30px">{items}</div>
  <div class="card" style="margin-top:20px;padding:19px 22px">
    <div class="eyebrow" style="color:var(--amber);margin-bottom:9px">Three bugs our own testing caught</div>
    <div style="font-size:14px;color:var(--ink-soft);line-height:1.6">
      <b style="color:var(--ink)">A splash screen that never hid</b> &mdash; a race between layout and the database left the app
      looking frozen and swallowing taps. &nbsp;&middot;&nbsp;
      <b style="color:var(--ink)">41 brand names that silently failed to resolve</b> &mdash; recognising a drug and having nothing
      to warn about are different things. &nbsp;&middot;&nbsp;
      <b style="color:var(--ink)">A swap chip that read &ldquo;Warfarin &rarr; Paracetamol&rdquo;</b> &mdash; now barred in code, with a test that
      stops any critical medicine ever being the one we suggest dropping.
    </div>
  </div>
</div>"""


SLIDES = [s01_title, s02_problem, s03_solution, s04_features, s05_demo,
          s06_how, s07_stack, s08_future, s09_close, a10_elderly, a11_testing]


# --------------------------------------------------------------------------- #

def find_chrome():
    for p in CHROME_CANDIDATES:
        if os.path.exists(p):
            return p
    found = shutil.which("chrome") or shutil.which("msedge")
    if found:
        return found
    sys.exit("Could not find Chrome or Edge to render the deck.")


def render(chrome, html, out_png):
    with tempfile.TemporaryDirectory() as tmp:
        page = os.path.join(tmp, "s.html")
        with open(page, "w", encoding="utf-8") as fh:
            fh.write(f'<!doctype html><meta charset="utf-8"><style>{CSS}</style>{html}')
        cmd = [chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars",
               "--allow-file-access-from-files",
               "--force-device-scale-factor=2",
               f"--user-data-dir={os.path.join(tmp, 'prof')}",
               f"--screenshot={out_png}", "--window-size=1280,720",
               "--virtual-time-budget=6000",
               "file:///" + page.replace("\\", "/")]
        subprocess.run(cmd, capture_output=True, timeout=180)
    return os.path.exists(out_png)


def main():
    chrome = find_chrome()
    shutil.rmtree(RENDER_DIR, ignore_errors=True)
    os.makedirs(RENDER_DIR, exist_ok=True)
    print(f"Rendering with: {chrome}")

    pngs = []
    for i, fn in enumerate(SLIDES, 1):
        out = os.path.join(RENDER_DIR, f"{i:02d}.png")
        ok = render(chrome, fn(), out)
        print(f"  slide {i:02d}  {fn.__name__:16} {'ok' if ok else 'FAILED'}")
        if not ok:
            sys.exit("A slide failed to render.")
        pngs.append(out)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for png in pngs:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(OUT_PPTX)

    print(f"\nSaved {OUT_PPTX}  ({len(pngs)} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
